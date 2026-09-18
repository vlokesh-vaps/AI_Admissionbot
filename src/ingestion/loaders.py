"""Load supported knowledge-base documents using LangChain document loaders."""

from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class LoadedDocument(Document):
    """A source document before chunking, fully compatible with LangChain Document."""

    def __init__(
        self,
        source_path: Path | str,
        title: str,
        text: str = "",
        metadata: dict[str, Any] | None = None,
        **kwargs: Any,
    ) -> None:
        path_obj = Path(source_path)
        meta = dict(metadata) if metadata else {}
        meta.setdefault("source", path_obj.name)
        meta.setdefault("source_path", str(path_obj))
        meta.setdefault("title", title)
        super().__init__(page_content=text, metadata=meta, **kwargs)

    @property
    def text(self) -> str:
        return self.page_content

    @property
    def source_path(self) -> Path:
        return Path(self.metadata.get("source_path", self.metadata.get("source", "")))

    @property
    def title(self) -> str:
        return str(self.metadata.get("title", ""))


def _clean_text(text: str) -> str:
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_pdf(path: Path) -> LoadedDocument:
    """Load a PDF using LangChain's PyPDFLoader, returned as a LoadedDocument."""
    loader = PyPDFLoader(str(path))
    pages = loader.load()
    page_texts: list[str] = []
    for i, page in enumerate(pages, start=1):
        content = _clean_text(page.page_content)
        if content:
            page_texts.append(f"[Page {i}]\n{content}")
    full_text = "\n\n".join(page_texts)
    if not full_text:
        raise ValueError(f"No extractable text found in {path.name}.")
    return LoadedDocument(
        source_path=path,
        title=path.stem,
        text=full_text,
    )


def load_docx(path: Path) -> LoadedDocument:
    """Load a DOCX file using python-docx, returned as a LoadedDocument."""
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise ImportError(
            "python-docx is required to load DOCX documents. "
            "Install with 'pip install python-docx'."
        ) from exc

    doc = DocxDocument(str(path))
    paragraphs = [_clean_text(p.text) for p in doc.paragraphs if _clean_text(p.text)]
    table_rows: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            cells = [_clean_text(cell.text) for cell in row.cells]
            table_rows.append(" | ".join(cell for cell in cells if cell))
    text = "\n".join(paragraphs + table_rows)
    if not text:
        raise ValueError(f"No extractable text found in {path.name}.")
    return LoadedDocument(
        source_path=path,
        title=path.stem,
        text=text,
    )


def load_document(path: Path) -> LoadedDocument:
    """Load one supported file, rejecting unsupported or empty documents."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}. Use PDF or DOCX.")
    return load_pdf(path) if suffix == ".pdf" else load_docx(path)


def load_documents(directory: Path) -> list[LoadedDocument]:
    """Load all supported files in deterministic filename order."""
    documents: list[LoadedDocument] = []
    for path in sorted(directory.iterdir()):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            documents.append(load_document(path))
    return documents
