import sys
from pathlib import Path
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.config import settings
from src.ingestion.loaders import load_document, load_documents


def test_docx_loader_extracts_paragraphs_and_tables() -> None:
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx is not installed in current environment")

    settings.documents_dir.mkdir(parents=True, exist_ok=True)
    fixture = settings.documents_dir / "_test_fixture.docx"
    document = Document()
    document.add_paragraph("Bachelor of Computer Applications")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Fee"
    table.rows[0].cells[1].text = "1000"
    document.save(fixture)
    try:
        loaded = load_document(fixture)
        assert "Bachelor of Computer Applications" in loaded.text
        assert "Fee | 1000" in loaded.text
        assert fixture in [item.source_path for item in load_documents(settings.documents_dir)]
    finally:
        fixture.unlink(missing_ok=True)
